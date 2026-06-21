#!/usr/bin/env python3
"""Build the Ricoman 'new website' leadership deck (PPTX) from the report content."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- Brand palette ----
INK     = RGBColor(0x0B, 0x0B, 0x0C)   # near-black
NAVY    = RGBColor(0x16, 0x33, 0x5C)   # deep brand navy
BLUE    = RGBColor(0x1D, 0x4E, 0xD8)   # accent blue
PAPER   = RGBColor(0xFF, 0xFF, 0xFF)
SOFT    = RGBColor(0xF4, 0xF6, 0xF8)   # light panel
MUTED   = RGBColor(0x5B, 0x62, 0x6D)
GREEN   = RGBColor(0x1A, 0x7F, 0x37)

prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def rect(s, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp

def txt(s, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """lines = list of (text, size, bold, color, space_after)"""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, (t, sz, b, c, sa) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sa); p.space_before = Pt(0)
        r = p.add_run(); r.text = t
        r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = c
        r.font.name = "Calibri"
    return tb

def bullets(s, x, y, w, h, items, size=16, color=INK, gap=10):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        # it = (head, body) or string
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.space_before = Pt(0); p.level = 0
        if isinstance(it, tuple):
            head, body = it
            r = p.add_run(); r.text = "▸ " + head
            r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = color; r.font.name = "Calibri"
            if body:
                r2 = p.add_run(); r2.text = "  " + body
                r2.font.size = Pt(size); r2.font.bold = False; r2.font.color.rgb = MUTED; r2.font.name = "Calibri"
        else:
            r = p.add_run(); r.text = "▸ " + it
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
    return tb

def header(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.35), PAPER)
    rect(s, 0, 0, Inches(0.22), Inches(1.35), BLUE)
    txt(s, Inches(0.6), Inches(0.18), Inches(12), Inches(1.1), [
        (kicker, 13, True, BLUE, 2),
        (title, 28, True, INK, 0),
    ])
    rect(s, Inches(0.6), Inches(1.2), Inches(12.1), Pt(2), SOFT)

def footer(s, n):
    txt(s, Inches(0.6), Inches(7.0), Inches(8), Inches(0.4),
        [("Ricoman — the new website", 10, False, MUTED, 0)])
    txt(s, Inches(11.8), Inches(7.0), Inches(1.0), Inches(0.4),
        [(str(n), 10, True, MUTED, 0)], align=PP_ALIGN.RIGHT)

# ============================================================ 1. TITLE
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, 0, SW, Inches(0.28), BLUE)
rect(s, 0, Inches(4.5), SW, Pt(2), RGBColor(0x2c,0x4a,0x78))
txt(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.2), [
    ("THE NEW RICOMAN WEBSITE", 40, True, PAPER, 8),
    ("Faster, cheaper to run, more reliable — and it wins business", 22, False, RGBColor(0xC9,0xD6,0xEC), 0),
])
txt(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(1.5), [
    ("A plain-English summary for the leadership report", 15, False, RGBColor(0x9F,0xB2,0xD4), 4),
    ("No technical knowledge needed", 13, False, RGBColor(0x7f,0x95,0xbf), 0),
])

# ============================================================ 2. IN ONE PARAGRAPH
s = slide(); header(s, "OVERVIEW", "In one sentence")
rect(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(3.2), SOFT)
rect(s, Inches(0.6), Inches(1.7), Inches(0.16), Inches(3.2), BLUE)
txt(s, Inches(1.1), Inches(1.95), Inches(11.2), Inches(2.7), [
    ("We haven't just repainted ricoman.com.", 24, True, INK, 12),
    ("We rebuilt it on modern foundations that are faster for customers, "
     "cheaper and more reliable to run, and far easier for the team to keep "
     "up to date — while quietly fixing years of hidden technical problems. "
     "And it does real commercial work: it captures sales leads, builds its "
     "own SEO, and lets customers design and specify products themselves.",
     19, False, NAVY, 0),
], anchor=MSO_ANCHOR.MIDDLE)
# three quick chips
chips = [("FASTER & CHEAPER", BLUE), ("MORE RELIABLE", GREEN), ("WINS BUSINESS", NAVY)]
cx = Inches(0.6)
for label, col in chips:
    rect(s, cx, Inches(5.3), Inches(3.9), Inches(0.95), col)
    txt(s, cx, Inches(5.3), Inches(3.9), Inches(0.95), [(label, 16, True, PAPER, 0)],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    cx += Inches(4.13)
footer(s, 2)

# ============================================================ 3. TECH DEBT
s = slide(); header(s, "1 · WHAT WE FIXED", "Years of hidden technical debt")
txt(s, Inches(0.6), Inches(1.55), Inches(12), Inches(0.5),
    [("Invisible from the front end — but costing money and putting the business at risk.", 15, False, MUTED, 0)])
# stat cards
stats = [
    ("75,911", "images in the old media library"),
    ("91%", "were exact duplicates (69,315)"),
    ("248 GB", "of bloat — backups failing since ~Oct 2025"),
]
cx = Inches(0.6)
for big, lab in stats:
    rect(s, cx, Inches(2.2), Inches(3.9), Inches(1.7), NAVY)
    txt(s, cx, Inches(2.35), Inches(3.9), Inches(1.0), [(big, 34, True, PAPER, 0)], align=PP_ALIGN.CENTER)
    txt(s, cx+Inches(0.2), Inches(3.25), Inches(3.5), Inches(0.6), [(lab, 12.5, False, RGBColor(0xC9,0xD6,0xEC), 0)], align=PP_ALIGN.CENTER)
    cx += Inches(4.13)
rect(s, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.2), SOFT)
rect(s, Inches(0.6), Inches(4.3), Inches(0.16), Inches(2.2), GREEN)
txt(s, Inches(1.0), Inches(4.5), Inches(11.4), Inches(0.5), [("The fix — and the payoff", 18, True, INK, 6)])
bullets(s, Inches(1.0), Inches(5.1), Inches(11.4), Inches(1.4), [
    ("Built-in Media Cleanup", "removed the duplicates safely."),
    ("Backups reliable again", "smaller, faster, and actually completing."),
    ("Lower, predictable hosting cost", "we store a fraction of the data."),
], size=15, gap=6)
footer(s, 3)

# ============================================================ 3b. BY THE NUMBERS
s = slide(); header(s, "BY THE NUMBERS", "Old site → new site, in real figures")
# big before/after metric tiles
def metric(s, x, y, w, old, arrow, new, label, newcol=GREEN):
    rect(s, x, y, w, Inches(1.95), SOFT)
    rect(s, x, y, w, Inches(0.42), INK)
    txt(s, x, y, w, Inches(0.42), [(label, 12, True, PAPER, 0)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    txt(s, x, y+Inches(0.5), w, Inches(0.5), [(old, 19, True, MUTED, 0)], align=PP_ALIGN.CENTER)
    txt(s, x, y+Inches(0.92), w, Inches(0.35), [(arrow, 15, True, BLUE, 0)], align=PP_ALIGN.CENTER)
    txt(s, x, y+Inches(1.22), w, Inches(0.6), [(new, 23, True, newcol, 0)], align=PP_ALIGN.CENTER)

metric(s, Inches(0.6),  Inches(1.75), Inches(3.9), "75,911 images", "▼  −91%", "6,596 unique", "MEDIA LIBRARY")
metric(s, Inches(4.73), Inches(1.75), Inches(3.9), "69,315 dupes", "removed", "0 duplicates", "DUPLICATE IMAGES")
metric(s, Inches(8.86), Inches(1.75), Inches(3.85), "248 GB, failing", "fixed", "Reliable backups", "HOSTING BACKUPS")

rect(s, Inches(0.6), Inches(4.1), Inches(12.1), Inches(2.4), PAPER, line=SOFT)
txt(s, Inches(0.95), Inches(4.25), Inches(11.4), Inches(0.5), [("What changed for the business", 17, True, INK, 6)])
bullets(s, Inches(0.95), Inches(4.8), Inches(11.6), Inches(1.6), [
    ("91% of the media library was waste", "— 69,315 of 75,911 images were exact duplicates from the old variant import."),
    ("Backups had been failing since ~Oct 2025", "— ballooned to 248 GB; cleaning the duplicates made them reliable again."),
    ("Lead capture: none → 5 routes", "— download gate, BIM request, newsletter, callback and saved-project packs."),
    ("Product edits: developer-only → self-serve", "— the team now edits spec tables and pages with no code."),
], size=13.5, gap=7)
txt(s, Inches(0.6), Inches(6.65), Inches(12), Inches(0.3),
    [("Figures from the media-cleanup audit of the migrated database. Live page-speed scores are measured at launch via the built-in SEO & Speed tool.", 9, False, MUTED, 0)])
footer(s, 4)

# ============================================================ 4. SELF-SERVE
s = slide(); header(s, "2 · EASIER TO RUN", "The team runs it themselves — no developer")
bullets(s, Inches(0.7), Inches(1.9), Inches(11.9), Inches(4.5), [
    ("Self-serve product pages.", "Edit the spec and variant tables (including which "
     "columns show) and drop content blocks into any product page — no code."),
    ("Editable landing pages.", "Casambi, Human Centric, Fire Safety, Made in Britain, "
     "Sustainability, Trade and more are all built from editable blocks."),
    ("One-click content tools", "for SEO copy and page layouts — they fill in starter "
     "content but never overwrite anything the team has written."),
], size=18, gap=18)
footer(s, 5)

# ============================================================ 5. SEO
s = slide(); header(s, "3 · IT MARKETS ITSELF", "Search visibility, built in")
bullets(s, Inches(0.7), Inches(1.8), Inches(11.9), Inches(4.8), [
    ("Automatic internal link-building.", "News and project stories auto-link relevant "
     "terms to the right product, category and sector pages."),
    ("Spec filtering for specifiers.", "Filter category pages by light output (lumens) "
     "and power (wattage) from the real product data."),
    ("Keyword-rich landing pages.", "Each category and key page has its own SEO title, "
     "description and intro/FAQ copy."),
    ("Rich Google results.", "Structured data for enhanced listings — business details, "
     "products, breadcrumbs and FAQs in search."),
    ("Built-in SEO health check.", "See at a glance which pages need attention."),
], size=16, gap=13)
footer(s, 6)

# ============================================================ 6. LEADS
s = slide(); header(s, "4 · IT CAPTURES SALES LEADS", "An active sales tool, not a brochure")
bullets(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(2.4), [
    ("Download gate.", "Document downloads ask for name, email and customer type — "
     "turning anonymous traffic into named leads."),
    ("Newsletter & request-a-callback", "forms feed the same pipeline."),
    ("Every lead in one place.", "A leads dashboard (type, contact, source, date) that "
     "can sync to a spreadsheet / CRM."),
    ("Saved projects & project packs", "— customers save lists and download a zipped "
     "pack of datasheets, instructions and photometric files."),
], size=15, gap=9)
# BIM highlight box
rect(s, Inches(0.6), Inches(5.45), Inches(12.1), Inches(1.25), NAVY)
rect(s, Inches(0.6), Inches(5.45), Inches(0.16), Inches(1.25), BLUE)
txt(s, Inches(1.0), Inches(5.55), Inches(11.5), Inches(1.1), [
    ('"Download BIM file" on every product page', 17, True, PAPER, 4),
    ("Specifiers drop BIM/Revit files straight into their building models — high-intent. "
     "Each request is logged as a lead and emailed to marketing with the exact product "
     "being specified. The moment a designer reaches for our product, marketing gets a warm lead.",
     13, False, RGBColor(0xC9,0xD6,0xEC), 0),
], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 7)

# ============================================================ 7. FLOW DESIGNER
s = slide(); header(s, "5 · CUSTOMERS DESIGN THEIR OWN", "The Flow+ Designer")
txt(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.6),
    [("A full-screen, in-browser tool to design a bespoke curved linear scheme — self-serve.", 15, False, MUTED, 0)])
bullets(s, Inches(0.7), Inches(2.1), Inches(11.9), Inches(3.2), [
    ("Draw the run.", "Drag in straights, curves and corners to any shape and custom "
     "length, with auto-snap so pieces join cleanly."),
    ("Design to scale on a real plan.", "Upload a CAD drawing / floor plan, set the scale, "
     "and lay the lighting out over the actual space."),
    ("Choose the spec.", "Colour temperature, DALI / Casambi dimming, drivers, RAL finish "
     "and emergency packs."),
    ("Instant 'Take Off'.", "A live parts list with the correct Ricoman order codes — so "
     "what they design is exactly what we quote and build."),
    ("Export & save", "as an image or take-off, and save straight to their project."),
], size=15, gap=10)
rect(s, Inches(0.6), Inches(5.75), Inches(12.1), Inches(1.0), SOFT)
rect(s, Inches(0.6), Inches(5.75), Inches(0.16), Inches(1.0), BLUE)
txt(s, Inches(1.0), Inches(5.75), Inches(11.4), Inches(1.0), [
    ("Why it matters:  bespoke quoting used to mean lots of back-and-forth. "
     "Now the customer arrives with a ready-to-quote, fully specified design — "
     "shorter sales cycle, fewer errors, a detailed high-intent enquiry.", 14, True, NAVY, 0),
], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 8)

# ============================================================ 8. FAST
s = slide(); header(s, "6 · PERFORMANCE", "Faster — and built to stay fast")
rect(s, Inches(0.6), Inches(2.0), Inches(12.1), Inches(3.0), SOFT)
txt(s, Inches(1.1), Inches(2.3), Inches(11.1), Inches(2.4), [
    ("Rendered natively and efficiently — no heavy page-builder overhead.", 19, True, INK, 12),
    ("Images are served in modern, lightweight formats, and each page loads only the "
     "code it needs. The result: a quicker experience for customers and better search "
     "rankings (Google factors in page speed).", 18, False, NAVY, 0),
], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 9)

# ============================================================ 9. ROADMAP
s = slide(); header(s, "ROADMAP", "Where we are")
steps = [
    ("1", "Recreate ricoman.com exactly on the new platform", "In progress — confirming each page matches live", BLUE),
    ("2", "Check & enhance SEO", "Foundations + audit tooling built", BLUE),
    ("3", "Launch — migrate the new site to live", "Runbook prepared", MUTED),
    ("4", "Feature & enhanced pages", "Built (Casambi, Fire Safety, Trade, …)", GREEN),
    ("5", "Reconnect live product data", "Final step, switched on last", MUTED),
]
y = Inches(1.8)
for num, title, sub, col in steps:
    rect(s, Inches(0.6), y, Inches(0.75), Inches(0.85), col)
    txt(s, Inches(0.6), y, Inches(0.75), Inches(0.85), [(num, 26, True, PAPER, 0)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    rect(s, Inches(1.45), y, Inches(11.25), Inches(0.85), SOFT)
    txt(s, Inches(1.7), y+Inches(0.08), Inches(10.8), Inches(0.75), [
        (title, 16, True, INK, 2),
        (sub, 12.5, False, MUTED, 0),
    ], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.0)
footer(s, 10)

# ============================================================ 10. BOTTOM LINE
s = slide(); header(s, "THE BOTTOM LINE", "Old site vs new site")
rows = [
    ("248 GB of bloat, 91% duplicate images", "Cleaned up, lean media library"),
    ("Backups failing since ~Oct 2025", "Reliable, fast backups"),
    ("Higher, growing hosting cost", "Lower, predictable hosting cost"),
    ("Changes needed a developer", "Team edits products & pages themselves"),
    ("A brochure", "Captures leads + builds its own SEO"),
    ("No BIM / specifier capture", '"Download BIM file" sends marketing a warm lead'),
    ("Bespoke quotes = lots of back-and-forth", "Flow+ Designer = ready-to-quote scheme"),
    ("Slower, page-builder overhead", "Fast, modern, search-friendly"),
]
# headers
rect(s, Inches(0.6), Inches(1.55), Inches(6.0), Inches(0.5), INK)
rect(s, Inches(6.7), Inches(1.55), Inches(6.0), Inches(0.5), BLUE)
txt(s, Inches(0.8), Inches(1.55), Inches(5.6), Inches(0.5), [("OLD SITE", 13, True, PAPER, 0)], anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(6.9), Inches(1.55), Inches(5.6), Inches(0.5), [("NEW SITE", 13, True, PAPER, 0)], anchor=MSO_ANCHOR.MIDDLE)
y = Inches(2.1)
for i, (old, new) in enumerate(rows):
    bg = PAPER if i % 2 == 0 else SOFT
    rect(s, Inches(0.6), y, Inches(6.0), Inches(0.56), bg)
    rect(s, Inches(6.7), y, Inches(6.0), Inches(0.56), bg)
    txt(s, Inches(0.8), y, Inches(5.7), Inches(0.56), [(old, 12, False, MUTED, 0)], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(6.9), y, Inches(5.7), Inches(0.56), [(new, 12, True, NAVY, 0)], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.6)
footer(s, 11)

# ============================================================ 11. CLOSE
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(2.6), SW, Pt(2), RGBColor(0x2c,0x4a,0x78))
txt(s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(2.6), [
    ("The new website is faster, cheaper to run, more reliable and easier to manage —", 22, True, PAPER, 10),
    ("and it actively works to win business, while having quietly fixed serious "
     "technical debt the old site was hiding.", 20, False, RGBColor(0xC9,0xD6,0xEC), 0),
], anchor=MSO_ANCHOR.MIDDLE)

out = "docs/Ricoman-New-Website.pptx"
prs.save(out)
print("saved", out, "with", len(prs.slides._sldIdLst), "slides")
